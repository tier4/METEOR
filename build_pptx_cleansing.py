#!/usr/bin/env python3
"""METEOR data-cleansing deck -> out/METEOR_data_cleansing.pptx.

Every filtering / cleansing mechanism between raw autolabels and the
training loss, with the measured effect where we have one."""
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Emu, Inches, Pt

SW, SH = Inches(13.333), Inches(7.5)
DARK = RGBColor(0x20, 0x28, 0x30)
ACC = RGBColor(0x0E, 0x6E, 0xB8)
GRAY = RGBColor(0x60, 0x68, 0x70)
GREEN = RGBColor(0x1B, 0x78, 0x37)
RED = RGBColor(0xC0, 0x39, 0x2B)
AMBER = RGBColor(0xB9, 0x77, 0x0E)
BG1 = RGBColor(0xF2, 0xF4, 0xF6)
prs = Presentation()
prs.slide_width, prs.slide_height = SW, SH
BLANK = prs.slide_layouts[6]


def slide(title=None, sub=None):
    s = prs.slides.add_slide(BLANK)
    if title:
        tb = s.shapes.add_textbox(Inches(0.5), Inches(0.22), Inches(12.3), Inches(0.7))
        p = tb.text_frame.paragraphs[0]
        p.text = title
        p.font.size = Pt(27)
        p.font.bold = True
        p.font.color.rgb = DARK
        ln = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(0.95),
                                Inches(12.3), Emu(1))
        ln.fill.solid()
        ln.fill.fore_color.rgb = ACC
        ln.line.fill.background()
        if sub:
            tb2 = s.shapes.add_textbox(Inches(0.5), Inches(1.0), Inches(12.3), Inches(0.4))
            p2 = tb2.text_frame.paragraphs[0]
            p2.text = sub
            p2.font.size = Pt(14)
            p2.font.italic = True
            p2.font.color.rgb = GRAY
    return s


def bullets(s, items, x=0.6, y=1.5, w=12.1, h=5.6, size=17):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    for i, it in enumerate(items):
        lvl, txt, col, sz = 0, "", None, None
        stack = list(it) if isinstance(it, tuple) else [it]
        while stack:
            e = stack.pop(0)
            if isinstance(e, RGBColor):
                col = e
            elif isinstance(e, str):
                txt = e
            elif isinstance(e, tuple):
                stack = list(e) + stack
            elif isinstance(e, int):
                if e in (0, 1):
                    lvl = e
                else:
                    sz = e
        base = size if lvl == 0 else size - 3
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = ("• " if lvl == 0 else "   – ") + txt
        p.font.size = Pt(sz if sz else base)
        p.font.color.rgb = col if col else DARK
        p.font.bold = sz is not None and lvl == 0
        p.space_after = Pt(6)


def big(s, lines, y=2.4):
    tb = s.shapes.add_textbox(Inches(0.8), Inches(y), Inches(11.7), Inches(3.0))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, (txt, sz, bold, col) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = txt
        p.font.size = Pt(sz)
        p.font.bold = bold
        p.font.color.rgb = col
        p.alignment = PP_ALIGN.CENTER
        p.space_after = Pt(10)


def table(s, rows, x, y, w, col_w, fs=13, hdr_fill=RGBColor(0xC9, 0xD9, 0xEA)):
    from pptx.util import Inches as In
    nr, nc = len(rows), len(rows[0])
    gt = s.shapes.add_table(nr, nc, In(x), In(y), In(w), In(0.34 * nr)).table
    for ci, cw in enumerate(col_w):
        gt.columns[ci].width = In(cw)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            c = gt.cell(ri, ci)
            c.text = str(val)
            pr = c.text_frame.paragraphs[0]
            pr.font.size = Pt(fs)
            pr.font.bold = ri == 0
            pr.font.color.rgb = DARK
            if ri == 0:
                c.fill.solid()
                c.fill.fore_color.rgb = hdr_fill
    return gt


def note(s, txt, y=6.85):
    tb = s.shapes.add_textbox(Inches(0.55), Inches(y), Inches(12.2), Inches(0.4))
    p = tb.text_frame.paragraphs[0]
    p.text = txt
    p.font.size = Pt(12)
    p.font.italic = True
    p.font.color.rgb = GRAY


# 1 ------------------------------------------------------------------
s = slide()
big(s, [
    ("METEOR", 54, True, ACC),
    ("Data Cleansing & Filtering in Training", 30, True, DARK),
    ("Every gate between raw autolabels and the loss — and what it measurably bought",
     17, False, GRAY),
    ("Principle: autolabels are noisy by construction. Never fake a missing label —", 15, False, DARK),
    ("mark it IGNORE and let the loss skip it. Filter at the weakest credible level.", 15, False, DARK),
], y=1.9)

# 2 ------------------------------------------------------------------
s = slide("The four filtering layers", "each label passes all four before it reaches a gradient")
table(s, [
    ["Layer", "Question it answers", "Mechanisms"],
    ["Scene", "is this recording usable?", "modality completeness, val-drive isolation, rolling list refresh"],
    ["Frame", "is this frame's GT trustworthy?", "LiDAR-coverage gates (gtcov), scene-head/tail trims, validity flags"],
    ["Label", "is this individual label real?", "camera confirmation, taxonomy drops, temporal median, ray carving"],
    ["Loss", "what if GT is absent anyway?", "255-ignore conventions, graph-preserving zero losses, per-class weights"],
], 0.7, 1.6, 12.0, [1.5, 3.6, 6.9], fs=14)
bullets(s, [
    ("Filtering is preferred over correction: a dropped label costs a little recall; "
     "a wrong label teaches the model to suppress its own confidence (measured on VRU detection).", 0),
], y=4.2, size=16)

# 3 ------------------------------------------------------------------
s = slide("Scene- and frame-level gates")
bullets(s, [
    ("LiDAR-coverage gate (gtcov): every frame's manifest stores [core, forward] "
     "coverage of the accumulated-LiDAR BEV GT; training drops frames below "
     "core 3% / forward 0.5% - thin GT teaches false negatives", 0),
    ("Scene head/tail trims (trim_start=3, trim_end): accumulation windows are "
     "one-sided at scene boundaries -> weakest BEV/occupancy GT simply cut", 0),
    ("Modality completeness: the round scene list requires the modalities the "
     "round trains on (e.g. r15+: agent_traj); 2,412/2,458 eligible scenes pass", 0),
    ("Validation isolation: one full held-out drive (all 273 sub-scenes of one "
     "recording day) - never appears in any training list", 0),
    ("E2E validity flag: no full 3 s pose future -> waypoint loss masked; "
     "instantaneous signals (v0/steer/accel) still written for EVERY frame "
     "(a scene-tail v0=0 once taught the model to 'hold' at 80 km/h)", 0),
    ("Temporal-pair validity: prev-frame images unreadable or pose missing -> "
     "prev BEV zeroed and flagged (pvalid=0), fusion sees an explicit blank", 0),
], size=15)

# 4 ------------------------------------------------------------------
s = slide("Label-level: camera confirmation of 3D boxes",
          "the single most impactful cleanser")
bullets(s, [
    ("Problem: LiDAR-annotation boxes include objects NO camera can see "
     "(full occlusion, distance) - unlearnable positives for a camera-only model", 0),
    ("Gate: project every 3D box into all 8 cameras; keep it only if a same-class "
     "2D autolabel box overlaps the projection (overlap/min-area > 0.3)", 0),
    ("Refinement (r15): VRUs have tiny 2D boxes - projection error dominates the "
     "overlap test. Class-dependent threshold 0.3 -> 0.15 for VRU", 0),
    ("Measured effect of the pair (relax + loss rebalance): VRU recall 0.08 -> 0.25, "
     "VRU share of GT boxes up to 38%", 0, GREEN),
    ("Same confirmed-instance set feeds 3D det, agent forecasting and the "
     "stationary flag -> one gate cleanses three tasks", 0),
], size=16)

# 5 ------------------------------------------------------------------
s = slide("Label-level: occupancy, segmentation, depth")
bullets(s, [
    ("Occupancy 'unknown' via ray carving: only voxels a LiDAR ray actually "
     "traversed (0.4 m steps) become 'free'; everything unobserved stays 255 "
     "- the model is never told an unseen voxel is empty", 0),
    ("Dynamic-object smear: statics accumulate over +-8 sweeps, but "
     "vehicles/pedestrians only from +-1 frame - kills motion ghosts in GT", 0),
    ("BEV lane 'marking' class dropped entirely as autolabel noise; sidewalk "
     "optionally trained as don't-care", 0),
    ("2D seg / depth: unlabeled pixels 255, missing modality -> all-255 frame "
     "(contributes nothing); mixed-resolution depth normalised, stale narrow-cam "
     "depth replaced by zeros rather than resampled", 0),
    ("2D boxes: degenerate (w<=0) boxes skipped; per-camera K normalised to a "
     "fixed KMAX=96 (mixed-shape GT once crashed a run mid-round - shapes are "
     "now normalised in the dataset, never in the files)", 0),
], size=15)

# 6 ------------------------------------------------------------------
s = slide("Label-level: temporal & semantic cleansing of TLR autolabels")
bullets(s, [
    ("Ego-relevance: front-NARROW telephoto first (sees only the travel "
     "corridor); front-WIDE accepted only in the central 50% of the image "
     "(side TLs belong to crossing roads)", 0),
    ("3-frame temporal median removes single-frame TLR colour flickers", 0),
    ("Largest-area (nearest) TL wins when several are confirmed", 0),
    ("Stationary flag GT: |GT displacement@3s| < 0.5 m only where the 3 s "
     "future exists - no guess at scene tails", 0),
    ("Risk-map GT inherits every upstream gate (camera-confirmed agents, "
     "carved occupancy) - cleansing composes", 0),
], size=16)

# 7 ------------------------------------------------------------------
s = slide("Loss-level: ignore-safe multi-task training",
          "missing GT must cost zero gradient - but keep the graph alive")
bullets(s, [
    ("Every aux loss returns pred.sum()*0 when its batch has no valid GT: "
     "zero-valued but graph-preserving - under DDP a truly unused head "
     "crashes the step ('did not receive grad'), a faked constant detaches it", 0),
    ("255-ignore in every dense loss (CE ignore_index, masked L1 at centres, "
     "tvalid masks per waypoint)", 0),
    ("This is what lets GT modalities STREAM IN mid-round: extraction runs "
     "concurrently with training; frames gain supervision as files appear", 0),
    ("Robust dataset: unreadable image -> neighbouring sample; concurrent "
     "manifest writes tolerated; per-scene calib tensors copied (shared-storage "
     "collate corruption)", 0),
], size=16)

# 8 ------------------------------------------------------------------
s = slide("QA loop: demos and metrics as cleansing detectors",
          "several gates above were DISCOVERED, not designed")
table(s, [
    ["Signal observed", "Root cause found", "Cleansing added"],
    ["E2E predicts 'hold' at 80 km/h", "scene-tail frames wrote v0=0", "instantaneous signals for every frame"],
    ["VRU recall stuck at 0.08", "confirmation gate too strict for small boxes", "class-dependent overlap threshold"],
    ["OCC hallucinates far statics", "unknown treated as free", "dense ray carving + 255 unknown"],
    ["OCC vehicles smeared", "dynamics accumulated over +-8 sweeps", "+-1-frame dynamics"],
    ["88 km/h 'drive' inside a garage", "multi-storey pose drift (t4dataset)", "visual QA gate on scene search; candidate lists screened"],
    ["3D det yaw noisy", "reg supervised only at exact centre cell", "3x3-neighbourhood targets"],
], 0.6, 1.6, 12.2, [3.6, 4.3, 4.3], fs=13)
note(s, "Rule: every regression demo/metric anomaly is treated as a potential GT defect first, model defect second.")

# 9 ------------------------------------------------------------------
s = slide("Selection & scheduling filters")
bullets(s, [
    ("Rolling rounds: the scene list is rebuilt every round from what the "
     "factory has finished - no frozen stale corpus", 0),
    ("--limit-train 46,000 samples/round: bounded epochs over a growing pool "
     "(fresh scenes rotate in each round)", 0),
    ("Warm starts with shape-filtered checkpoint loading: new heads initialise, "
     "everything else continues - no re-learning from scratch after GT fixes", 0),
    ("Best-checkpoint selection by composite score (mIoU - 0.01*min(ADEc,5)): "
     "a single-metric 'best' once resurrected a pre-fix regression", 0),
    ("Loss-level class weights double as soft cleansing: far-range positives "
     "damped where 768x432 input physically cannot resolve the object "
     "(unlearnable labels suppress confidence everywhere if left at full weight)", 0),
], size=16)

# 10 -----------------------------------------------------------------
s = slide("Summary: gate -> measured effect")
table(s, [
    ["Gate", "Effect"],
    ["gtcov coverage + trims", "removes ~8% weakest frames; road-edge recall stable across rounds"],
    ["camera confirmation (+VRU relax)", "VRU recall 0.08 -> 0.25; GT VRU share 38%"],
    ["ray-carved unknown (occupancy)", "far-static hallucination eliminated in demo QA"],
    ["+-1-frame dynamics (occupancy)", "moving-object smear removed from GT"],
    ["instantaneous E2E signals", "high-speed 'hold' failure eliminated"],
    ["TLR median + ego-relevance", "stable whole-image TL labels (acc 0.71 after 1 epoch)"],
    ["far-positive damping", "veh near-corridor recall 0.66-0.70 with P recovery at 0.45 thr (0.65->0.76)"],
    ["ignore-safe streaming", "GT extraction runs concurrently with training, zero crashes since r7"],
], 1.0, 1.5, 11.3, [4.6, 6.7], fs=13)
note(s, "All numbers from the held-out validation drive; details in docs/TRAINING.md rolling-rounds table.")

import os
os.makedirs("out", exist_ok=True)
prs.save("out/METEOR_data_cleansing.pptx")
print("saved out/METEOR_data_cleansing.pptx")
