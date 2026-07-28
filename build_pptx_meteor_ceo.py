#!/usr/bin/env python3
"""METEOR CEO introduction deck (Japanese, ~15 min) -> out/METEOR_CEO.pptx.

First-time executive introduction. Selling points: Zero Human Code, Zero Human
Annotation, No Map, depth-based vehicle-robust perception, lightweight edge
deployment, self-guardrail.

All block diagrams are built from NATIVE PowerPoint shapes + connectors (no
image insertion) so they can be edited afterwards.
"""
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

SW, SH = Inches(13.333), Inches(7.5)
C_NAVY = RGBColor(0x15, 0x22, 0x38); C_BLUE = RGBColor(0x1E, 0x6E, 0xB8)
C_GRAY = RGBColor(0x60, 0x68, 0x70); C_GREEN = RGBColor(0x1B, 0x78, 0x37)
C_RED = RGBColor(0xC0, 0x39, 0x2B); C_AMBER = RGBColor(0xE0, 0x8A, 0x1E)
C_WHITE = RGBColor(0xFF, 0xFF, 0xFF); C_TEAL = RGBColor(0x0E, 0x9A, 0xA0)
C_EDGE = RGBColor(0x4A, 0x55, 0x60)
F_GRAY = RGBColor(0xEE, 0xEE, 0xEE); F_BLUE = RGBColor(0xCF, 0xE2, 0xF3)
F_AMBER = RGBColor(0xFD, 0xE7, 0xB5); F_GREEN = RGBColor(0xD9, 0xEA, 0xD3)
F_PURP = RGBColor(0xE8, 0xD5, 0xF2)
prs = Presentation(); prs.slide_width, prs.slide_height = SW, SH
BLANK = prs.slide_layouts[6]


def _fill(shape, rgb):
    shape.fill.solid(); shape.fill.fore_color.rgb = rgb
    shape.line.fill.background()


def slide(title=None, sub=None, band=C_BLUE):
    s = prs.slides.add_slide(BLANK)
    if title:
        tb = s.shapes.add_textbox(Inches(0.55), Inches(0.25), Inches(12.2),
                                  Inches(0.8))
        p = tb.text_frame.paragraphs[0]
        p.text = title; p.font.size = Pt(28); p.font.bold = True
        p.font.color.rgb = C_NAVY
        ln = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.55), Inches(1.02),
                                Inches(12.2), Pt(3))
        _fill(ln, band)
        if sub:
            tb2 = s.shapes.add_textbox(Inches(0.55), Inches(1.06), Inches(12.2),
                                       Inches(0.45))
            p2 = tb2.text_frame.paragraphs[0]
            p2.text = sub; p2.font.size = Pt(14); p2.font.italic = True
            p2.font.color.rgb = C_GRAY
    return s


def bullets(s, items, x=0.7, y=1.6, w=11.9, h=5.5, size=17):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.auto_size = MSO_AUTO_SIZE.NONE
    for i, it in enumerate(items):
        lvl, txt, col = 0, it, None
        if isinstance(it, tuple):
            lvl, txt = it[0], it[1]; col = it[2] if len(it) > 2 else None
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = ("■ " if lvl == 0 else "     ・") + txt
        p.font.size = Pt(size if lvl == 0 else size - 3)
        p.font.color.rgb = col or C_NAVY
        p.font.bold = (lvl == 0 and col is not None)
        p.space_after = Pt(7)


# ---- native block-diagram helpers (editable shapes + connectors) ----
def nbox(s, x, y, w, h, title, sub="", fc=F_BLUE, fs=12.5, sfs=9, tc=C_NAVY):
    b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y),
                           Inches(w), Inches(h))
    b.fill.solid(); b.fill.fore_color.rgb = fc
    b.line.color.rgb = C_EDGE; b.line.width = Pt(1.1)
    b.shadow.inherit = False
    tf = b.text_frame; tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    p = tf.paragraphs[0]; p.text = title; p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(fs); p.font.bold = True; p.font.color.rgb = tc
    if sub:
        p2 = tf.add_paragraph(); p2.text = sub; p2.alignment = PP_ALIGN.CENTER
        p2.font.size = Pt(sfs); p2.font.color.rgb = RGBColor(0x55, 0x5C, 0x63)
    return b


def _arrow(conn, col):
    ln = conn.line._get_or_add_ln()
    tail = ln.makeelement(qn("a:tailEnd"),
                          {"type": "triangle", "w": "med", "len": "med"})
    ln.append(tail)
    conn.line.color.rgb = col; conn.line.width = Pt(2.0)


def link(s, a, b, col=C_BLUE, side="h"):
    """Connector glued between shapes a->b. side 'h' = a.right->b.left,
    'v' = a.bottom->b.top. Stays attached when boxes are moved (editable)."""
    if side == "h":
        x1, y1 = a.left + a.width, a.top + a.height // 2
        x2, y2 = b.left, b.top + b.height // 2
        pa, pb = 3, 1
    else:
        x1, y1 = a.left + a.width // 2, a.top + a.height
        x2, y2 = b.left + b.width // 2, b.top
        pa, pb = 2, 0
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    _arrow(c, col)
    try:
        c.begin_connect(a, pa); c.end_connect(b, pb)
    except Exception:
        pass
    return c


def caption(s, txt, y=6.9, col=C_AMBER, size=12):
    tb = s.shapes.add_textbox(Inches(0.6), Inches(y), Inches(12.1), Inches(0.4))
    p = tb.text_frame.paragraphs[0]; p.text = txt; p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(size); p.font.italic = True; p.font.color.rgb = col


def table(s, rows, x, y, w, col_w, fs=13, head_fill=RGBColor(0x1E, 0x6E, 0xB8),
          zebra=RGBColor(0xF2, 0xF4, 0xF6)):
    nr, nc = len(rows), len(rows[0])
    t = s.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w),
                           Inches(0.5 * nr)).table
    for ci, cw in enumerate(col_w):
        t.columns[ci].width = Inches(cw)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            c = t.cell(ri, ci); c.text = str(val)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            pr = c.text_frame.paragraphs[0]
            pr.font.size = Pt(fs); pr.font.bold = (ri == 0 or ci == 0)
            if ri == 0:
                c.fill.solid(); c.fill.fore_color.rgb = head_fill
                pr.font.color.rgb = C_WHITE
            else:
                c.fill.solid()
                c.fill.fore_color.rgb = zebra if ri % 2 else C_WHITE
                pr.font.color.rgb = (C_RED if ci == 1 else
                                     C_GREEN if ci == 2 else C_NAVY)
    return t


def bignum(s, items, y=2.2):
    n = len(items); w = 12.0 / n
    for i, (num, lab) in enumerate(items):
        x = 0.7 + i * w
        card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x),
                                  Inches(y), Inches(w - 0.35), Inches(2.4))
        _fill(card, RGBColor(0xF2, 0xF4, 0xF6))
        t = card.text_frame; t.word_wrap = True
        t.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = t.paragraphs[0]; p.text = num; p.alignment = PP_ALIGN.CENTER
        p.font.size = Pt(34); p.font.bold = True; p.font.color.rgb = C_BLUE
        p2 = t.add_paragraph(); p2.text = lab; p2.alignment = PP_ALIGN.CENTER
        p2.font.size = Pt(13); p2.font.color.rgb = C_NAVY


def demo_slide(title, sub=None, band=C_BLUE):
    s = slide(title, sub, band=band)
    frame = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(2.3),
                               Inches(1.9), Inches(8.7), Inches(4.9))
    frame.fill.solid(); frame.fill.fore_color.rgb = RGBColor(0x0E, 0x14, 0x22)
    frame.line.color.rgb = C_BLUE; frame.line.width = Pt(2)
    tf = frame.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.text = "▶"; p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(60); p.font.color.rgb = C_WHITE
    p2 = tf.add_paragraph(); p2.text = "（デモ動画をここに挿入）"
    p2.alignment = PP_ALIGN.CENTER; p2.font.size = Pt(18)
    p2.font.color.rgb = RGBColor(0xB9, 0xCF, 0xE8)
    return s


# ---------------------------------------------------------------- 1 title
s = prs.slides.add_slide(BLANK)
bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH); _fill(bg, C_NAVY)
tb = s.shapes.add_textbox(Inches(1.0), Inches(2.3), Inches(11.3), Inches(2))
p = tb.text_frame.paragraphs[0]
p.text = "METEOR"; p.font.size = Pt(66); p.font.bold = True
p.font.color.rgb = C_WHITE; p.alignment = PP_ALIGN.CENTER
tb2 = s.shapes.add_textbox(Inches(1.0), Inches(3.9), Inches(11.3), Inches(1.4))
for i, (t, sz, c) in enumerate([
        ("人手ゼロで作る、自動運転の「眼」と「脳」", 26, C_WHITE),
        ("ラベル無し・地図無し・人手コード無し ── カメラだけで走る", 16,
         RGBColor(0xB9, 0xCF, 0xE8)),
        ("", 8, C_WHITE),
        ("学習データ: 日本全国 9,600+シーン ＝ 走行約80時間（まる3.3日分）を人手ゼロで学習",
         14, C_TEAL)]):
    p = tb2.text_frame.paragraphs[0] if i == 0 else tb2.text_frame.add_paragraph()
    p.text = t; p.font.size = Pt(sz); p.font.color.rgb = c
    p.alignment = PP_ALIGN.CENTER; p.font.bold = (i == 0)

# ---------------------------------------------------------------- 2 exec summary
s = slide("エグゼクティブ・サマリー", "3行で言うと")
bullets(s, [
    (0, "METEORは、人手のアノテーションもコード記述も無しに、走行データだけから"
        "自動運転の認識・計画AIを自律的に作り続けるエンジンです。", C_NAVY),
    (0, "カメラのみ（深度をAIが推定）で、地図に依存せず、軽量にエッジ配備。"
        "決定論的な自己ガードレールで安全に動きます。", C_NAVY),
    (0, "現時点で人手ラベル0枚のまま9,600シーン（走行約80時間＝3.3日分・約125万フレーム）を自動学習し、"
        "車両検出・車線・信号・経路計画までを1モデルで出力しています。", C_GREEN),
], size=18)
bignum(s, [("0", "人手ラベル"), ("0", "人手コード行"), ("0", "HDマップ"),
           ("80h", "走行データ(3.3日分)")], y=4.6)

# ---------------------------------------------------------------- 2.5 glossary
s = slide("はじめに ── この資料で使う3つの言葉", "これだけ分かれば全部読めます")
gl = [("BEV（俯瞰図）", "クルマの真上から見た地図のような絵。"
       "8台のカメラ映像をAIが1枚の俯瞰図に変換し、その上で道路や他車を理解します。", F_BLUE),
      ("アノテーション", "AIに見せる「正解ラベル」を人が手作業で付ける仕事。"
       "従来は数万枚規模で必要。METEORはこれがゼロ（自動生成）。", F_AMBER),
      ("E2E（経路計画）", "カメラ映像から「次にどう走るか」の線まで一気通貫にAIが出すこと。"
       "METEORは認識と計画を1つのモデルで同時に行います。", F_GREEN)]
for i, (t, d, fc) in enumerate(gl):
    card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                              Inches(0.7 + i * 4.2), Inches(2.1),
                              Inches(3.9), Inches(3.6))
    card.fill.solid(); card.fill.fore_color.rgb = fc
    card.line.color.rgb = C_EDGE; card.line.width = Pt(1.2)
    card.shadow.inherit = False
    tf = card.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.15)
    p = tf.paragraphs[0]; p.text = t; p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(19); p.font.bold = True; p.font.color.rgb = C_NAVY
    p2 = tf.add_paragraph(); p2.text = "\n" + d
    p2.font.size = Pt(13); p2.font.color.rgb = C_NAVY

# ---------------------------------------------------------------- 3 problem
s = slide("なぜ自動運転の「認識」は高コストなのか", "従来アプローチの3つの重荷",
          band=C_RED)
bullets(s, [
    (0, "アノテーション地獄", C_RED),
    (1, "数万〜数十万枚の画像を人手でラベル付け。1都市分でも数ヶ月・高コスト。"),
    (0, "HDマップ依存", C_RED),
    (1, "事前に精密地図を作成・維持。地図の無い場所は走れない。"),
    (0, "重い開発と重いハード", C_RED),
    (1, "大規模な人手開発、LiDAR等の高価センサー、データセンター級の計算機。"),
    (0, "→ 結果、スケール（新地域・新車種への展開）が遅く、コストが青天井。", C_NAVY),
])

# ---------------------------------------------------------------- 4 what is meteor
s = slide("METEORとは ── 自律的に成長する認識エンジン")
bullets(s, [
    (0, "CoMET（コメット）＝ 走行データに「正解ラベル」を自動で付ける仕組み"
        "（Co-MLOpsプロジェクトの自動ラベリング基盤）。いわばラベルの自動工場。",
     C_AMBER),
    (0, "METEOR（メテオ）＝ CoMETが作ったラベル付きデータで育つ、"
        "自動運転の認識・計画AI。学習→自己改善→配備まで人手を介さず自動で回します。",
     C_GREEN),
], y=1.4, size=15)
# native pipeline diagram
py = 3.7; ph = 1.7
b1 = nbox(s, 0.55, py, 2.2, ph, "走行データ", "DRS収集・日本全国\n8カメラ+LiDAR", fc=F_GRAY)
b2 = nbox(s, 3.15, py, 2.35, ph, "CoMET＝ラベル自動工場", "正解ラベルを自動付与\n(Co-MLOps基盤・人手0)", fc=F_AMBER, fs=11)
b3 = nbox(s, 5.9, py, 2.05, ph, "自己学習", "マルチタスク\nBEVモデル", fc=F_BLUE)
b4 = nbox(s, 8.35, py, 2.0, ph, "自己改善", "Refinerが\n出力を自動補正", fc=F_PURP)
b5 = nbox(s, 10.75, py, 2.0, ph, "エッジ配備", "TensorRT/C++\nカメラのみ", fc=F_GREEN)
for a, b in ((b1, b2), (b2, b3), (b3, b4), (b4, b5)):
    link(s, a, b)
caption(s, "収集 → ラベル → 学習 → 改善 → 配備 の全工程を自動化"
        "（人手のラベリングもコード記述も不要）", y=5.8)

# ---------------------------------------------------------------- 4.5 dev method
s = slide("開発のやり方 ── 人間は「指示」、AIが「開発」",
          "コンセプト・機能は人間が指示。残りの全工程はAIが主導（Zero Human Code）")


def person_icon(sl, cx, cy, h, col):
    """simple editable person pictogram (head + shoulders)"""
    hd = h * 0.36
    head = sl.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - hd / 2), Inches(cy),
                               Inches(hd), Inches(hd))
    _fill(head, col)
    bw, bh = h * 0.62, h * 0.5
    body = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                               Inches(cx - bw / 2), Inches(cy + hd + 0.04),
                               Inches(bw), Inches(bh))
    _fill(body, col)


def robot_icon(sl, cx, cy, h, col):
    """simple editable robot pictogram (antenna + head + eyes + mouth)"""
    tip = sl.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - h * 0.07),
                              Inches(cy - h * 0.02), Inches(h * 0.14),
                              Inches(h * 0.14))
    _fill(tip, col)
    ant = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(cx - 0.015),
                              Inches(cy + h * 0.1), Inches(0.03),
                              Inches(h * 0.18))
    _fill(ant, col)
    hw = h * 0.92
    head = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                               Inches(cx - hw / 2), Inches(cy + h * 0.26),
                               Inches(hw), Inches(h * 0.72))
    _fill(head, col)
    for dx in (-hw * 0.22, hw * 0.22):
        eye = sl.shapes.add_shape(MSO_SHAPE.OVAL,
                                  Inches(cx + dx - h * 0.075),
                                  Inches(cy + h * 0.44), Inches(h * 0.15),
                                  Inches(h * 0.15))
        _fill(eye, C_WHITE)
    mo = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                             Inches(cx - hw * 0.26), Inches(cy + h * 0.76),
                             Inches(hw * 0.52), Inches(h * 0.09))
    _fill(mo, C_WHITE)


# --- human zone (left, amber) ---
hb = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.45),
                        Inches(1.8), Inches(2.5), Inches(3.55))
hb.fill.solid(); hb.fill.fore_color.rgb = F_AMBER
hb.line.color.rgb = C_AMBER; hb.line.width = Pt(1.75); hb.shadow.inherit = False
person_icon(s, 1.7, 2.0, 1.0, C_AMBER)
ht = s.shapes.add_textbox(Inches(0.5), Inches(3.6), Inches(2.4), Inches(0.95))
p = ht.text_frame.paragraphs[0]; p.text = "人間"; p.alignment = PP_ALIGN.CENTER
p.font.size = Pt(19); p.font.bold = True; p.font.color.rgb = C_NAVY
p2 = ht.text_frame.add_paragraph()
p2.text = "コンセプト・機能を\n指示するだけ"; p2.alignment = PP_ALIGN.CENTER
p2.font.size = Pt(11.5); p2.font.color.rgb = C_NAVY
qb = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.62),
                        Inches(4.55), Inches(2.16), Inches(0.62))
qb.fill.solid(); qb.fill.fore_color.rgb = C_WHITE
qb.line.color.rgb = C_AMBER; qb.line.width = Pt(1); qb.shadow.inherit = False
p = qb.text_frame.paragraphs[0]; p.text = "「◯◯できる機能が欲しい」"
p.alignment = PP_ALIGN.CENTER; p.font.size = Pt(9.5); p.font.italic = True
p.font.color.rgb = C_NAVY

# --- AI zone (right band, blue) ---
band = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(3.9),
                          Inches(1.8), Inches(8.9), Inches(3.55))
band.fill.solid(); band.fill.fore_color.rgb = RGBColor(0xF2, 0xF6, 0xFB)
band.line.color.rgb = C_BLUE; band.line.width = Pt(1.75)
band.shadow.inherit = False
robot_icon(s, 7.05, 1.95, 0.62, C_BLUE)
bh_ = s.shapes.add_textbox(Inches(7.45), Inches(2.05), Inches(3.6),
                           Inches(0.45))
p = bh_.text_frame.paragraphs[0]; p.text = "AIが主導（Claude Fable 5）"
p.font.size = Pt(15); p.font.bold = True; p.font.color.rgb = C_BLUE
# 5 steps, each with its own editable icon shape above the box
steps = [(MSO_SHAPE.GEAR_6, "モデル開発", "設計・学習・評価"),
         (MSO_SHAPE.CAN, "データセット\n読み込み", "新データ自動取込"),
         (MSO_SHAPE.LEFT_RIGHT_ARROW, "データ変換", "学習形式へ自動変換"),
         (MSO_SHAPE.FUNNEL, "データ\nクレンジング", "ノイズ自動除去"),
         (MSO_SHAPE.CUBE, "モデル管理", "配備・自動復旧")]
sb = []
for i, (icon, t, sub) in enumerate(steps):
    bx = 4.1 + i * 1.75
    if icon == MSO_SHAPE.LEFT_RIGHT_ARROW:      # wide so it reads as an arrow
        ic = s.shapes.add_shape(icon, Inches(bx + 0.44), Inches(2.8),
                                Inches(0.68), Inches(0.32))
    else:
        ic = s.shapes.add_shape(icon, Inches(bx + 0.55), Inches(2.72),
                                Inches(0.45), Inches(0.45))
    _fill(ic, C_BLUE)
    sb.append(nbox(s, bx, 3.3, 1.55, 1.3, t, sub, fc=F_BLUE,
                   fs=11, sfs=8.5))
for a, b in zip(sb, sb[1:]):
    link(s, a, b)
# 24h self-improvement loop inside the AI band
lo = s.shapes.add_shape(MSO_SHAPE.DONUT, Inches(5.45), Inches(4.72),
                        Inches(0.5), Inches(0.5))
_fill(lo, C_TEAL)
lt = s.shapes.add_textbox(Inches(6.1), Inches(4.78), Inches(6.0), Inches(0.4))
p = lt.text_frame.paragraphs[0]
p.text = "24時間 自己改善ループ（クラッシュ復旧・再学習も自動）"
p.font.size = Pt(11.5); p.font.bold = True; p.font.color.rgb = C_TEAL

# --- human <-> AI arrows ---
c1 = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, hb.left + hb.width,
                            Inches(2.9), band.left, Inches(2.9))
_arrow(c1, C_AMBER)
c2 = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, band.left, Inches(4.35),
                            hb.left + hb.width, Inches(4.35))
_arrow(c2, C_TEAL)
for lx, ly, txt, col in ((2.87, 2.5, "指示", C_AMBER),
                         (2.87, 4.45, "報告・提案", C_TEAL)):
    lb = s.shapes.add_textbox(Inches(lx), Inches(ly), Inches(1.1), Inches(0.35))
    p = lb.text_frame.paragraphs[0]; p.text = txt; p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(10.5); p.font.bold = True; p.font.color.rgb = col

# --- workload ratio bar (who does the work) ---
rt = s.shapes.add_textbox(Inches(0.45), Inches(5.62), Inches(2.0), Inches(0.35))
p = rt.text_frame.paragraphs[0]; p.text = "作業量の割合:"
p.font.size = Pt(12); p.font.bold = True; p.font.color.rgb = C_NAVY
seg_h = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.95), Inches(5.95),
                           Inches(0.55), Inches(0.45))
_fill(seg_h, C_AMBER)
seg_a = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(2.5), Inches(5.95),
                           Inches(10.35), Inches(0.45))
_fill(seg_a, C_BLUE)
p = seg_a.text_frame.paragraphs[0]
p.text = "AI ～95%（モデル開発・データ処理・運用のすべて）"
p.alignment = PP_ALIGN.CENTER; p.font.size = Pt(12); p.font.bold = True
p.font.color.rgb = C_WHITE
hl = s.shapes.add_textbox(Inches(1.35), Inches(6.45), Inches(1.9),
                          Inches(0.35))
p = hl.text_frame.paragraphs[0]; p.text = "人間 ～5%（指示）"
p.font.size = Pt(10.5); p.font.bold = True; p.font.color.rgb = C_AMBER
caption(s, "人間は「何を作るか」を決めるだけ ── "
        "モデル開発からデータ処理・運用まで、作り方はすべてAIが自律実行",
        y=6.95, col=C_GREEN, size=13)

# ---------------------------------------------------------------- 10.5 dev history
s = slide("開発の歩み ── 約3週間で多タスク統合まで", "全コードAI記述による開発スピード")
# native horizontal timeline (editable shapes)
tl_y = 3.7
line = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.2), Inches(tl_y),
                          Inches(10.9), Pt(4))
_fill(line, C_BLUE)
mile = [
    (1.7, "7/10", "開発スタート", "BEV Segmentation・\n3D BBox 推論", C_BLUE, True),
    (6.3, "7/17", "機能追加", "E2E（経路計画）\n機能を追加", C_TEAL, False),
    (11.3, "現在", "統合・自己改善", "多タスク統合・TensorRT実装\n・自己改善ループ稼働", C_GREEN, True),
]
for x, date, head, desc, col, above in mile:
    dot = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x - 0.16), Inches(tl_y - 0.13),
                             Inches(0.34), Inches(0.34))
    _fill(dot, col)
    # date label on the line
    db = s.shapes.add_textbox(Inches(x - 0.7), Inches(tl_y + 0.25), Inches(1.4),
                              Inches(0.4))
    dp = db.text_frame.paragraphs[0]; dp.text = date; dp.alignment = PP_ALIGN.CENTER
    dp.font.size = Pt(16); dp.font.bold = True; dp.font.color.rgb = col
    # description card above or below
    cy = tl_y - 1.85 if above else tl_y + 0.75
    card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x - 1.5),
                              Inches(cy), Inches(3.0), Inches(1.5))
    _fill(card, RGBColor(0xF2, 0xF4, 0xF6))
    tf = card.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    hp = tf.paragraphs[0]; hp.text = head; hp.alignment = PP_ALIGN.CENTER
    hp.font.size = Pt(14); hp.font.bold = True; hp.font.color.rgb = col
    dp2 = tf.add_paragraph(); dp2.text = desc; dp2.alignment = PP_ALIGN.CENTER
    dp2.font.size = Pt(11); dp2.font.color.rgb = C_NAVY
caption(s, "人手を待たずに24時間開発を継続 ── わずか約3週間で認識から計画まで1モデルに統合",
        y=6.4, col=C_GREEN, size=14)

# ---------------------------------------------------------------- 5 model architecture
s = slide("モデルアーキテクチャ（53Mパラメータ・CNN構成）",
          "8カメラ → 深度付き俯瞰変換 → 時系列融合 → 10タスク同時出力 → 補正・安全層")
# --- inputs column ---
i1 = nbox(s, 0.3, 1.55, 1.75, 1.05, "8カメラ", "サラウンド6 +\n前後望遠2 (432×768)",
          fc=F_GRAY, fs=11, sfs=8)
i2 = nbox(s, 0.3, 2.8, 1.75, 0.8, "車速・運動履歴", "v0 + 直近2.8秒", fc=F_GRAY,
          fs=9.5, sfs=8)
i3 = nbox(s, 0.3, 3.8, 1.75, 0.8, "Driving Command", "直進/左折/右折(任意)",
          fc=F_GRAY, fs=9, sfs=8)
i4 = nbox(s, 0.3, 4.8, 1.75, 0.8, "LiDAR(任意)", "学習時のみ使用可",
          fc=F_GRAY, fs=9.5, sfs=8)
# --- image branch ---
bb = nbox(s, 2.5, 1.55, 1.85, 1.5, "CNNバックボーン\n+FPN", "カメラ毎に共有重み",
          fc=F_BLUE, fs=10.5, sfs=8)
d2 = nbox(s, 2.5, 3.35, 1.85, 0.95, "2D出力", "21クラスセグ\n10クラス2D箱",
          fc=F_BLUE, fs=10, sfs=8)
dp = nbox(s, 4.75, 1.55, 1.85, 0.95, "深度分布推定", "ピクセル毎の距離確率",
          fc=F_AMBER, fs=10.5, sfs=8)
ipm = nbox(s, 4.75, 2.75, 1.85, 1.15, "深度ゲート\nIPM投影", "特徴×深度を俯瞰へ",
           fc=F_AMBER, fs=10.5, sfs=8)
# --- BEV branch ---
bev = nbox(s, 7.0, 1.55, 1.9, 1.0, "BEVグリッド", "800×500・0.2m/セル\n前後±80m 左右±50m",
           fc=F_GREEN, fs=10.5, sfs=8)
tmp = nbox(s, 7.0, 2.8, 1.9, 1.1, "時系列メモリ融合", "0.4/1.2/2.8秒前のBEVを\n自車移動分ずらして重ねる",
           fc=F_GREEN, fs=10, sfs=8)
enc = nbox(s, 7.0, 4.15, 1.9, 0.9, "BEVエンコーダ", "256ch 共有特徴",
           fc=F_GREEN, fs=10.5, sfs=8)
# --- heads grid (2 cols x 5) ---
heads = [("BEVセグ", "車線・停止線等 9クラス"), ("3D物体検出", "車両/歩行者+向き・速度"),
         ("信号状態", "青/黄/赤 + 自車関連性"), ("E2E経路計画", "3候補・車速/Command入力"),
         ("他者予測", "3秒先の軌跡場"), ("リスク地図", "衝突危険度の分布"),
         ("3D占有(OCC)", "立体ボクセル"), ("未知障害物", "落下物等・密マップ"),
         ("車線グラフ", "ベクトル接続関係"), ("占有フロー", "動き場")]
hx0, hy0 = 9.35, 1.5
hboxes = []
for i, (t, d) in enumerate(heads):
    r, c = divmod(i, 2)
    hb2 = nbox(s, hx0 + c * 1.72, hy0 + r * 0.78, 1.62, 0.68, t, d,
               fc=F_BLUE, fs=9, sfs=6.5)
    hboxes.append(hb2)
# --- refiner + guardrail ---
rf = nbox(s, 9.35, hy0 + 5 * 0.78 + 0.15, 3.34, 0.62, "Refiner（タスク毎の残差補正）",
          "凍結出力に補正のみ加算＝劣化しない構造", fc=F_PURP, fs=9.5, sfs=7)
gd = nbox(s, 9.35, hy0 + 5 * 0.78 + 0.92, 3.34, 0.62, "自己ガードレール（決定論）",
          "衝突・赤信号・逸脱チェック → 不合格なら安全停止", fc=F_PURP, fs=9.5, sfs=7)
# --- wires ---
link(s, i1, bb)
link(s, bb, dp); link(s, bb, ipm)
c_ = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, bb.left + bb.width // 2,
                            bb.top + bb.height, d2.left + d2.width // 2, d2.top)
_arrow(c_, C_BLUE)
link(s, dp, bev); link(s, ipm, tmp)
link(s, bev, tmp, side="v", col=C_GREEN)
link(s, tmp, enc, side="v", col=C_GREEN)
link(s, enc, hboxes[8])
c2_ = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, enc.left + enc.width,
                             enc.top + enc.height // 2, Inches(hx0),
                             Inches(hy0 + 0.35))
_arrow(c2_, C_BLUE)
c3_ = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                             Inches(hx0 + 1.6), Inches(hy0 + 5 * 0.78),
                             rf.left + rf.width // 2, rf.top)
_arrow(c3_, RGBColor(0x8E, 0x44, 0xAD))
link(s, rf, gd, side="v", col=C_TEAL)
# aux inputs: thin ELBOW wires to their consumers (E2E head / IPM) --
# straight diagonals crossed the whole diagram and were unreadable
e2e_head = hboxes[3]
for src, col in ((i2, C_TEAL), (i3, C_TEAL)):
    ce = s.shapes.add_connector(MSO_CONNECTOR.ELBOW,
                                src.left + src.width,
                                src.top + src.height // 2,
                                e2e_head.left + e2e_head.width // 2,
                                e2e_head.top + e2e_head.height)
    ce.line.color.rgb = col; ce.line.width = Pt(1.1)
    try:
        ce.begin_connect(src, 3); ce.end_connect(e2e_head, 2)
    except Exception:
        pass
cl = s.shapes.add_connector(MSO_CONNECTOR.ELBOW, i4.left + i4.width,
                            i4.top + i4.height // 2, ipm.left + ipm.width // 2,
                            ipm.top + ipm.height)
cl.line.color.rgb = C_GRAY; cl.line.width = Pt(1.1)
try:
    cl.begin_connect(i4, 3); cl.end_connect(ipm, 2)
except Exception:
    pass
caption(s, "CNN構成＝現行SoC（Orin）で動作可能な現実的な設計（LLM/世界モデル化は現行車載HWでは困難なため未着手）"
        "／ 学習データはCoMET自動ラベル（人手0）", y=6.95, col=C_GRAY, size=10.5)


# ------------------------------------------------- 5b extended architecture
def _dashed(shape):
    ln = shape.line._get_or_add_ln()
    d = ln.makeelement(qn("a:prstDash"), {"val": "dash"})
    ln.append(d)


s = slide("拡張アーキテクチャ ── 信号認識・地図情報の追加（開発中）",
          "どちらも「あっても無くても」動く任意入力 ── 無い時は従来と完全に同一の出力")
# --- new optional inputs (dashed = optional/new) ---
ntl = nbox(s, 0.3, 1.6, 2.3, 1.25, "信号認識モジュール（外部）",
           "前方カメラの2D枠 +\n色（赤/黄/青）+ 矢印方向", fc=F_AMBER, fs=10.5, sfs=8.5)
nsd = nbox(s, 0.3, 3.15, 2.3, 1.45, "SDマップ（無料OSM）",
           "道路網・交差点・横断歩道\n・法定速度・一時停止\n（将来HDマップに差替可）",
           fc=F_AMBER, fs=10.5, sfs=8.5)
nex = nbox(s, 0.3, 4.9, 2.3, 0.9, "従来入力", "8カメラ・車速履歴\nCommand・LiDAR(任意)",
           fc=F_GRAY, fs=10, sfs=8.5)
_dashed(ntl); _dashed(nsd)
# --- adapters ---
atl = nbox(s, 3.15, 1.6, 2.35, 1.25, "信号状態ベクトル",
           "自車関連の信号を選択\n状態+距離を数値化", fc=F_PURP, fs=10.5, sfs=8.5)
asd = nbox(s, 3.15, 3.15, 2.35, 1.45, "地図ラスタ変換 + Stem",
           "自車周辺±80mを俯瞰画像化\nゼロ初期化CNN（無入力=無影響）",
           fc=F_PURP, fs=10.5, sfs=8.5)
_dashed(atl); _dashed(asd)
# --- existing pipeline (condensed) ---
core = nbox(s, 6.1, 3.0, 2.5, 1.75, "既存パイプライン",
            "バックボーン → 深度付き\n俯瞰変換 → 時系列融合\n→ BEVエンコーダ",
            fc=F_GREEN, fs=11, sfs=9)
# --- heads ---
h_e2e = nbox(s, 9.3, 1.7, 3.3, 0.95, "E2E経路計画（強化）",
             "赤信号で停止・青で発進\n交差点の先の道路を地図で先読み",
             fc=F_BLUE, fs=10.5, sfs=8.5)
h_tls = nbox(s, 9.3, 2.9, 3.3, 0.8, "信号状態ヘッド（強化）",
             "認識モジュールとの融合で高信頼化", fc=F_BLUE, fs=10.5, sfs=8.5)
h_rest = nbox(s, 9.3, 3.95, 3.3, 0.8, "BEVセグ・3D検出 ほか8タスク",
              "遠方・遮蔽部分の精度向上", fc=F_BLUE, fs=10.5, sfs=8.5)
gd2 = nbox(s, 9.3, 5.0, 3.3, 0.8, "自己ガードレール（決定論）",
           "赤信号通過・逸脱・衝突をチェック", fc=F_PURP, fs=10.5, sfs=8.5)
# --- wires ---
link(s, ntl, atl, col=C_AMBER)
link(s, nsd, asd, col=C_AMBER)
link(s, atl, h_e2e, col=C_AMBER)
link(s, asd, core, col=C_AMBER)
link(s, nex, core, col=C_GRAY)
link(s, core, h_e2e); link(s, core, h_tls); link(s, core, h_rest)
link(s, h_rest, gd2, side="v", col=C_TEAL)
ce2 = s.shapes.add_connector(MSO_CONNECTOR.ELBOW,
                             atl.left + atl.width, atl.top + atl.height // 2,
                             h_tls.left, h_tls.top + h_tls.height // 2)
ce2.line.color.rgb = C_AMBER; ce2.line.width = Pt(1.3)
try:
    ce2.begin_connect(atl, 3); ce2.end_connect(h_tls, 1)
except Exception:
    pass
caption(s, "点線＝新規の任意入力（実データ検証済み：信号の色・矢印は既存ラベルで供給可、"
        "標識・横断歩道は無料OSMから取得済み）", y=6.6, col=C_AMBER, size=11)
caption(s, "期待効果：交差点の先の道路をカメラ視界外でも予測（地図prior）／"
        "信号・標識に整合した停止・発進の計画", y=7.0, col=C_GRAY, size=10.5)

# ---------------------------------------------------------------- 6 screen guide (before the video)
s = slide("デモ画面の見方", "1枚の画面に「見る・測る・理解する・決める」が全部出ます")
img_w, img_h = 8.6, 4.84
pic = s.shapes.add_picture("docs/media/ceo_demo_frame.png", Inches(0.45),
                           Inches(1.7), Inches(img_w), Inches(img_h))
guide = [
    ("(1) 8カメラ+認識結果", "車両・歩行者・信号を映像上に表示",
     C_BLUE, 2.0, 0.45 + img_w * 0.35, 1.75 + img_h * 0.12),
    ("(2) 距離の予測(深度)", "色=距離。カメラだけで測距",
     C_AMBER, 3.15, 0.45 + img_w * 0.35, 1.75 + img_h * 0.55),
    ("(3) 立体空間の理解", "ボクセル=立体の占有マップ",
     C_GREEN, 4.3, 0.45 + img_w * 0.12, 1.75 + img_h * 0.85),
    ("(4) 俯瞰図(BEV)+走行計画", "緑線=AIの走行計画。GUARD OK=\n決定論の安全チェック合格",
     C_TEAL, 5.45, 0.45 + img_w * 0.88, 1.75 + img_h * 0.5),
]
for t, d, col, ly, ax, ay in guide:
    lb = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(9.35),
                            Inches(ly), Inches(3.6), Inches(1.0))
    lb.fill.solid(); lb.fill.fore_color.rgb = C_WHITE
    lb.line.color.rgb = col; lb.line.width = Pt(2); lb.shadow.inherit = False
    tf = lb.text_frame; tf.word_wrap = True
    tf.margin_left = Inches(0.08); tf.margin_top = Inches(0.03)
    p = tf.paragraphs[0]; p.text = t
    p.font.size = Pt(13.5); p.font.bold = True; p.font.color.rgb = col
    p2 = tf.add_paragraph(); p2.text = d
    p2.font.size = Pt(10.5); p2.font.color.rgb = C_NAVY
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(9.35),
                               Inches(ly + 0.5), Inches(ax), Inches(ay))
    _arrow(c, col)
caption(s, "人が運転中に頭の中でやっていること（見る・距離感・立体把握・進路決め）を、"
        "1つのAIが毎フレーム実行", y=6.85, col=C_GREEN, size=13)

# ---------------------------------------------------------------- 6.5 METEOR demo
demo_slide("METEOR デモ ── 実走行での認識・計画",
           "カメラのみ・地図なしで、BEV・3D物体・信号・経路計画をリアルタイム出力")

# ---------------------------------------------------------------- 6.7 self-improvement
s = slide("自動で賢くなる ── 人手ゼロの改善ループ実績",
          "ラウンド（自動学習の周回）毎に、人が何もしなくても精度が上がる")
imp = [("BEVセグ精度 (mIoU)", "0.324", "0.334", "↑", C_GREEN),
       ("車両検出の再現率", "47%", "51%", "↑", C_GREEN),
       ("経路計画の誤差", "0.43m", "0.39m", "↓ (小さいほど良い)", C_BLUE)]
for i, (t, a, b, ar, col) in enumerate(imp):
    card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                              Inches(0.7 + i * 4.2), Inches(2.0),
                              Inches(3.9), Inches(2.9))
    _fill(card, RGBColor(0xF2, 0xF4, 0xF6))
    tf = card.text_frame; tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.text = t; p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(16); p.font.bold = True; p.font.color.rgb = C_NAVY
    p2 = tf.add_paragraph(); p2.text = f"{a} → {b}"
    p2.alignment = PP_ALIGN.CENTER; p2.font.size = Pt(30); p2.font.bold = True
    p2.font.color.rgb = col
    p3 = tf.add_paragraph(); p3.text = ar
    p3.alignment = PP_ALIGN.CENTER; p3.font.size = Pt(13); p3.font.color.rgb = C_GRAY
bullets(s, [
    (0, "改善のサイクル（データ追加 → 自動ラベル → 学習 → 評価 → 弱点の自動補正）を"
        "24時間、AIが自律で回し続けた結果です。", C_NAVY),
    (0, "→ 精度向上が人件費ではなく「計算時間」で買える構造。", C_AMBER),
], y=5.3, size=15)

# ---------------------------------------------------------------- 7 differentiators
s = slide("METEORの6つの強み（＝競争優位）", "この組み合わせが他に無い")
strengths = [
    ("① Zero Human Annotation", "人手ラベル0枚。CoMETが正解ラベルを自動生成し、"
     "9,600+シーン/125万フレームをラベル費ゼロで学習。新データ追加も数時間。",
     C_GREEN, F_GREEN),
    ("② Zero Human Code", "モデル設計から配備まで全コードをAIが記述。"
     "24時間自律で開発・復旧・改善。3週間で認識→計画まで統合。",
     C_GREEN, F_GREEN),
    ("③ No Map", "HDマップ不要。その場のカメラ映像だけで走行空間を再構成。"
     "地図の作成・維持費ゼロ、未整備地域や道路変化にも即対応。",
     C_GREEN, F_GREEN),
    ("④ Depthベースで車両ロバスト", "ピクセル毎の深度推定で正確な俯瞰変換。"
     "LiDARは学習時の教師のみ、推論はカメラだけ（センサーコスト低減）。",
     C_GREEN, F_GREEN),
    ("⑤ 軽量・エッジ実装", "TensorRT/C++変換済み・実測108ms（L40S）。"
     "CNN構成で現行SoC（Orin）動作可。J6等の車両構成へ最適化予定。",
     C_TEAL, F_BLUE),
    ("⑥ 自己ガードレール", "AIの計画を決定論チェックが常時監視"
     "（衝突・赤信号・逸脱）。不合格なら安全停止。構造で安全を担保。",
     C_TEAL, F_BLUE),
]
for i, (t, d, tc, fc) in enumerate(strengths):
    r, c = divmod(i, 3)
    card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                              Inches(0.55 + c * 4.15), Inches(1.65 + r * 2.65),
                              Inches(3.95), Inches(2.45))
    card.fill.solid(); card.fill.fore_color.rgb = fc
    card.line.color.rgb = tc; card.line.width = Pt(1.5)
    card.shadow.inherit = False
    tf = card.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.12)
    tf.margin_top = Inches(0.08)
    p = tf.paragraphs[0]; p.text = t
    p.font.size = Pt(15); p.font.bold = True; p.font.color.rgb = tc
    p2 = tf.add_paragraph(); p2.text = d
    p2.font.size = Pt(11); p2.font.color.rgb = C_NAVY
caption(s, "各項目の詳細データ・比較は次頁の対比表とホワイトペーパーへ",
        y=6.95, col=C_GRAY, size=11)

# ---------------------------------------------------------------- 8 vs conventional (native table)
s = slide("従来手法との違い（一目で）")
table(s, [
    ["観点", "従来の自動運転認識", "METEOR"],
    ["アノテーション", "数万〜数十万枚を人手（数ヶ月）", "人手ゼロ・自動生成（数時間で拡張）"],
    ["HDマップ", "事前作成・維持が必要", "地図不要（どこでも走行）"],
    ["開発コード", "大規模な人手開発", "全コードをAIが記述"],
    ["センサー", "高価なLiDAR依存も多い", "カメラのみで推論（深度をAI推定）"],
    ["安全担保", "モデル任せになりがち", "決定論的な自己ガードレール"],
    ["配備", "重く高価な計算機", "軽量・エッジ実装"],
], 0.9, 1.5, 11.5, [2.6, 4.6, 4.3], fs=14)

# ---------------------------------------------------------------- 8.5 vs VAD
s = slide("METEOR vs VAD ── 最先端E2E研究との違い",
          "VAD: 学界の代表的なカメラE2E手法（ベクトル化シーン表現）")
# two pipeline rows as native shapes
vy, vh = 1.5, 0.8
lb = s.shapes.add_textbox(Inches(0.5), Inches(vy + 0.2), Inches(1.35),
                          Inches(0.45))
p = lb.text_frame.paragraphs[0]; p.text = "VAD"; p.font.size = Pt(17)
p.font.bold = True; p.font.color.rgb = C_RED
v1 = nbox(s, 1.9, vy, 1.9, vh, "カメラ映像", fc=F_GRAY, fs=11.5)
v2 = nbox(s, 4.15, vy, 2.15, vh, "Transformer BEV", "重い注意機構", fc=F_GRAY,
          fs=11.5, sfs=8.5)
v3 = nbox(s, 6.65, vy, 2.3, vh, "ベクトル化シーン", "地図・他車", fc=F_GRAY,
          fs=11.5, sfs=8.5)
v4 = nbox(s, 9.3, vy, 1.75, vh, "経路計画", fc=F_GRAY, fs=11.5)
for a, b in ((v1, v2), (v2, v3), (v3, v4)):
    link(s, a, b, col=C_RED)
vt = nbox(s, 11.4, vy, 1.5, vh, "人手ラベル\nで学習", fc=RGBColor(0xF6, 0xD0, 0xCC),
          fs=10, sfs=8)
my, mh = 2.55, 0.8
lb = s.shapes.add_textbox(Inches(0.5), Inches(my + 0.2), Inches(1.4),
                          Inches(0.45))
p = lb.text_frame.paragraphs[0]; p.text = "METEOR"; p.font.size = Pt(15)
p.font.bold = True; p.font.color.rgb = C_GREEN
m1 = nbox(s, 1.9, my, 1.9, mh, "8カメラ 360°", fc=F_BLUE, fs=11.5)
m2 = nbox(s, 4.15, my, 2.15, mh, "深度ベースBEV", "CNN・軽量", fc=F_BLUE,
          fs=11.5, sfs=8.5)
m3 = nbox(s, 6.65, my, 2.3, mh, "多タスク認識+計画", "セグ/3D/信号/占有…",
          fc=F_BLUE, fs=10.5, sfs=8)
m4 = nbox(s, 9.3, my, 1.75, mh, "ガードレール", "決定論安全層", fc=F_PURP,
          fs=10.5, sfs=8)
for a, b in ((m1, m2), (m2, m3), (m3, m4)):
    link(s, a, b, col=C_GREEN)
mt = nbox(s, 11.4, my, 1.5, mh, "自動ラベル\nで学習(人手0)", fc=F_GREEN,
          fs=10, sfs=8)
# can / cannot table
table(s, [
    ["観点", "VAD（E2E研究の代表）", "METEOR"],
    ["E2E経路計画・3D検出・車線", "○", "○（同じ土俵）"],
    ["信号・深度・占有・未知障害物", "−（対象外）", "○（1モデルで出力）"],
    ["学習ラベル", "人手アノテーションが必要", "人手ゼロ（CoMET自動ラベル）"],
    ["開発", "研究者が実装", "AIが全コード記述"],
    ["安全担保", "NN出力に依存", "決定論ガードレールで二重化"],
    ["計算・実装", "Transformer・研究実装", "CNN・車載SoC（Orin）動作可"],
    ["公開ベンチ実績", "○（nuScenes）", "今後（White Paperで公開予定）"],
], 0.9, 3.65, 11.5, [3.3, 4.1, 4.1], fs=12.5)

# ---------------------------------------------------------------- 15 capabilities
s = slide("何ができるか ── 1モデルで多タスク", "カメラのみ・地図なしの実測値（検証データ）")
bullets(s, [
    (0, "BEVセグメンテーション（車線・停止線・横断歩道・走行領域）: mIoU 0.334"),
    (0, "3D物体検出：車両 適合率 ~80%・再現率 51%（遠方40m超も強化中）"),
    (0, "信号認識：全体 74%（青信号 90%+）"),
    (0, "経路計画(E2E)：カーブ誤差 0.39m（直線含む平均 0.87m）"),
    (0, "加えて：深度・占有・車線グラフ・未知障害物・他者挙動予測・リスク地図"),
    (1, "後段のRefinerが、遠方や細部を自動で補正しさらに高精度化。"),
    (0, "未学習の米国道路でも車両・信号・経路の基本認識が動作（ゼロショット汎化）。",
     C_GREEN),
], size=16)

# ---------------------------------------------------------------- 16 business value
s = slide("事業価値", "コスト構造を根本から変える", band=C_AMBER)
bullets(s, [
    (0, "アノテーション費 → 自動ラベリング基盤（CoMET）の強化用のみ。走行データ量に比例して増えない（データを増やしてもラベル費ゼロ）。"),
    (0, "地図作成・維持費 → 不要（No Map）。"),
    (0, "開発人件費 → 大幅圧縮（コードをAIが記述、24時間稼働）。"),
    (0, "センサーBOM → 低減（カメラ中心、LiDAR任意）。"),
    (0, "→ 新地域・新車種への展開が「データを入れるだけ」でスケール。"
        "限界費用が劇的に低い。", C_AMBER),
])

# ---------------------------------------------------------------- 17 roadmap
s = slide("今後の予定（ロードマップ）", band=C_AMBER)
bullets(s, [
    (0, "データのロバスト化", C_AMBER),
    (1, "NVIDIA Cosmosで生成した多様・希少シーンのデータで学習し、実環境への"
        "堅牢性を高める。"),
    (0, "エッジ最適化・車載SoC実装", C_AMBER),
    (1, "必要な機能に絞り込んだ上で最適化を実施。"),
    (1, "J6（ミニバス）のセンサー構成に合わせたファインチューニング"
        "（7カメラ構成への絞り込み）。"),
    (1, "NVIDIA SoC「Orin」および Renesas「R-Car Gen5」への実装を進める。"),
    (0, "リファレンスAI化とWhite Paper公開", C_AMBER),
    (1, "Reference AI（オープンソース）として公開し、業界標準・エコシステム形成を狙う。"),
    (1, "技術詳細をWhite Paperとして公開："
        "github.com/tier4/METEOR/blob/main/paper/main.pdf"),
    (0, "並行して：遠方・未知障害物の精度向上、車種・地域の横展開、自己学習ループの"
        "完全自律化。", C_NAVY),
], size=15)

# ---------------------------------------------------------------- 18 Cosmos demo
demo_slide("Cosmos生成データによるロバスト化 ── デモ",
           "NVIDIA Cosmosで生成した多様・希少シーンでの認識・計画の頑健性",
           band=C_AMBER)

# ---------------------------------------------------------------- 19 closing
s = prs.slides.add_slide(BLANK)
bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH); _fill(bg, C_NAVY)
tb = s.shapes.add_textbox(Inches(1.0), Inches(2.3), Inches(11.3), Inches(3.2))
lines = [("人手ゼロで、自動運転の認識を作り続ける。", 30, C_WHITE, True),
         ("ラベル無し・地図無し・コード無し。カメラだけで、どこでも走る。",
          18, RGBColor(0xB9, 0xCF, 0xE8), False),
         ("", 10, C_WHITE, False),
         ("METEOR ── スケールする自律認識エンジン", 22, C_TEAL, True),
         ("", 12, C_WHITE, False),
         ("Repository:  github.com/tier4/METEOR", 15,
          RGBColor(0xB9, 0xCF, 0xE8), False),
         ("White Paper:  github.com/tier4/METEOR/blob/main/paper/main.pdf", 14,
          RGBColor(0x9F, 0xB8, 0xD4), False)]
for i, (t, sz, c, b) in enumerate(lines):
    p = tb.text_frame.paragraphs[0] if i == 0 else tb.text_frame.add_paragraph()
    p.text = t; p.font.size = Pt(sz); p.font.color.rgb = c
    p.font.bold = b; p.alignment = PP_ALIGN.CENTER

out = "out/METEOR_CEO.pptx"
prs.save(out)
print("saved", out, "-", len(prs.slides._sldIdLst), "slides")
